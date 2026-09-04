from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AskPDF_Implementation_and_Code_Guide.docx"
PURPLE = "B44627"
NAVY = "24231F"
MUTED = "706D64"
PALE = "F2EEE4"
GREEN = "68704B"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table) -> None:
    table_pr = table._tbl.tblPr
    table_borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9DCEC")
        table_borders.append(edge)
    table_pr.append(table_borders)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths=None,
    *,
    trailing_space: bool = True,
    font_size: float = 8,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    repeat_header(table.rows[0])
    keep_row_together(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, PURPLE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        for index, value in enumerate(values):
            if row_index % 2:
                shade(cells[index], "F8F8FC")
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(NAVY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    if trailing_space:
        doc.add_paragraph()
    return table


def add_code(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shade_node = OxmlElement("w:shd")
    shade_node.set(qn("w:fill"), "F3F4F8")
    p_pr.append(shade_node)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_callout(doc: Document, title: str, text: str, color: str = PALE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    keep_row_together(table.rows[0])
    cell = table.cell(0, 0)
    shade(cell, color)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    lead = paragraph.add_run(f"{title}: ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(PURPLE)
    paragraph.add_run(text)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.2)
styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name, size, color in (
    ("Title", 34, NAVY),
    ("Subtitle", 13, MUTED),
    ("Heading 1", 21, NAVY),
    ("Heading 2", 14, PURPLE),
    ("Heading 3", 11, NAVY),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("AskPDF · Implementation & Code Guide · ")
footer_run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Cover
brand = doc.add_paragraph()
brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
brand.paragraph_format.space_before = Pt(68)
run = brand.add_run("ASK")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor.from_string(NAVY)
run = brand.add_run("PDF")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor.from_string(PURPLE)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(12)
title_run = title.add_run("AskPDF Implementation\nand Code Guide")
title_run.font.name = "Aptos Display"
title_run.font.size = Pt(34)
title_run.font.color.rgb = RGBColor.from_string(NAVY)
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Citation First PDF Question Answering Production TypeScript Monorepo")
doc.add_paragraph()
add_callout(
    doc,
    "Build contract",
    "PDF work is asynchronous; every accepted factual answer is tied to authorized document versions and exact PDF pages; tenant scope is enforced in storage, queues, retrieval, cache keys, and database filters.",
)
image_path = ROOT / "docs" / "design" / "askpdf-editorial-workspace-concept.png"
if image_path.exists():
    picture = doc.add_picture(str(image_path), width=Inches(6.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph("Implementation baseline · 5 September 2026")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.runs[0].italic = True
caption.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
doc.add_page_break()

heading(doc, "1. What is being built")
doc.add_paragraph(
    "AskPDF is a multi-tenant RAG web application for students, researchers, analysts, and teams who need answers they can verify. Users must sign in with email or E.164 phone number plus password, upload PDFs directly to private object storage, monitor background processing, select ready sources, ask questions, and open the page behind every citation."
)
add_table(
    doc,
    ["User problem", "Implemented capability", "Where it appears"],
    [
        ["Large PDFs are slow to inspect", "Page-aware background extraction and hybrid retrieval", "Upload status and Ask PDFs workspace"],
        ["AI answers can invent facts", "Evidence threshold, structured claims, citation-to-chunk validation, abstention", "Answer cards and citation drawer"],
        ["Sensitive documents can cross boundaries", "Workspace authorization inside every resource and retrieval query", "Invisible security boundary; safe 401/403/404 responses"],
        ["Long processing blocks web requests", "Direct upload plus BullMQ worker stages", "Progress states remain responsive"],
        ["Research needs source inspection", "Short-lived signed source URL with exact page number", "Citation click opens source drawer"],
    ],
)
add_callout(doc, "MVP boundary", "Quiz, study notes, debate mode, flashcards, spaced repetition, and collaborative learning remain planned v2 work. Their designs are recorded, but no unfinished buttons or fake server behavior are shipped in the MVP.")

heading(doc, "2. System shape")
add_code(
    doc,
    "Browser (React/Vite)\n"
    "  ├── HTTPS JSON + SSE ──> Express API ──> MongoDB / Redis / MinIO / Gemini\n"
    "  └── signed PUT ───────────────────────> MinIO or S3\n"
    "                                                │\n"
    "                                 BullMQ job ───┴──> Worker\n"
    "                                      parse → normalize → chunk → embed → publish progress",
)
add_bullets(
    doc,
    [
        "The web application owns presentation, session state, direct upload, retry affordances, streamed progress, and citation navigation.",
        "The API owns validation, authentication, authorization, metadata, signed URLs, conversation orchestration, retrieval, citation validation, and public error contracts.",
        "The worker owns all expensive PDF and embedding work. It limits temporary bytes, streams object downloads, verifies checksums, checkpoints stages, and cleans up even after failure.",
        "Shared packages point inward: contracts and configuration have no application dependencies; infrastructure adapters depend on contracts; applications compose adapters and business services.",
    ],
)

heading(doc, "3. Build it from scratch: exact order")
build_rows = [
    ["1", "Repository", "Pin Node/pnpm; create strict TS workspaces, formatting, lint, Vitest, Playwright.", "Every later change has reproducible quality gates."],
    ["2", "Configuration", "Parse every environment variable with Zod and fail startup on invalid security or limit values.", "Broken deployment config cannot become a runtime surprise."],
    ["3", "Infrastructure", "Start Mongo replica set, authenticated Redis, and private MinIO bucket with health checks.", "Transactions, queues, and object storage are available locally."],
    ["4", "Contracts/models", "Define request/job/event schemas and tenant-owned Mongoose records with indexes.", "API and worker exchange versioned, validated shapes."],
    ["5", "Authentication", "Implement Argon2id registration/login, cookie sessions, refresh rotation, reuse detection, CSRF, membership roles.", "All product routes are protected."],
    ["6", "Upload vertical slice", "Create upload intent → signed PUT → complete → enqueue deterministic ingestion job.", "Browser never proxies PDF bytes through Express."],
    ["7", "Ingestion", "Stream download, validate signature/hash/limits, extract pages, normalize, chunk, embed, commit ready version.", "Documents become queryable without blocking the API."],
    ["8", "Retrieval", "Embed question, run tenant-filtered vector and lexical searches, fuse ranks, enforce evidence floor.", "Only authorized, sufficient evidence enters generation."],
    ["9", "Generation", "Ask Gemini for a strict claims schema; validate pages, chunks, excerpts, and version; persist answer/citations.", "Unsupported claims never reach the user."],
    ["10", "Streaming", "Publish processing progress and answer lifecycle via SSE with keepalive and cleanup.", "Long work stays observable and cancellable."],
    ["11", "Frontend", "Build compulsory auth, documents, upload, collection, chat, citations, all states, responsive theme.", "The end-to-end product is usable on desktop and mobile."],
    ["12", "Hardening", "Add readiness probes, structured logs, quotas/rate limits, secret scanning, malicious-file tests.", "Operational and abuse controls guard production."],
    ["13", "Evaluation", "Version a labeled PDF/question dataset and calculate retrieval, faithfulness, citation, refusal metrics.", "Quality changes are measurable and reproducible."],
    ["14", "Release", "Run CI gates, staging smoke/load tests, backup/rollback rehearsal, deploy web/API/worker independently.", "Release is reversible and evidence-backed."],
]
add_table(doc, ["Order", "Foundation", "Minute implementation sequence", "Exit condition"], build_rows)

heading(doc, "4. Repository map: file-by-file responsibility")
module_rows = [
    ["packages/contracts/src/index.ts", "Boundary contracts", "Strict Zod schemas, constants, inferred TS types", "All external inputs and model outputs"],
    ["packages/config/src/index.ts", "Configuration", "Typed environment parsing and cross-field validation", "Startup and configurable limits"],
    ["packages/database/src/models/index.ts", "Persistence", "Tenant-aware documents, runs, chunks, chat, citations, indexes", "Durable application state"],
    ["packages/storage/src/index.ts", "Object storage", "Signed PUT/GET, metadata, byte streams, readiness", "Direct upload and source viewing"],
    ["packages/queue/src/index.ts", "Jobs/events", "Deterministic BullMQ jobs and Redis progress bus", "Asynchronous orchestration"],
    ["packages/rag/src/index.ts", "PDF/RAG primitives", "Extraction, normalization, chunks, fusion, citation checks", "Evidence production and ranking"],
    ["packages/ai/src/index.ts", "Model boundary", "Gemini embeddings and structured cited-claim generation", "Provider calls with typed output"],
    ["apps/api/src/application/auth-service.ts", "Identity", "Register/login/refresh/logout/profile and CSRF", "Compulsory access control"],
    ["apps/api/src/application/resource-service.ts", "Use cases", "Uploads, documents, collections, conversations, retrieval, citations", "Core product behavior"],
    ["apps/api/src/app.ts", "HTTP transport", "Routes, middleware, SSE, error envelope, security headers", "Public API contract"],
    ["apps/worker/src/worker.ts", "Background compute", "Retry-safe staged ingestion and deletion", "Ready document versions"],
    ["apps/web/src/api.ts", "Browser gateway", "Credentialed fetch, CSRF refresh, signed upload", "Typed UI/server integration"],
    ["apps/web/src/pages/*.tsx", "Screens", "Auth, documents, collection, chat, citations", "Visible user journeys"],
    ["apps/web/src/styles.css", "Design system", "Editorial surfaces, light and dark tokens, restrained motion, responsiveness", "Requested visual identity"],
]
add_table(doc, ["File", "Topic", "What the block does", "Final-site effect"], module_rows)

heading(doc, "5. Foundational contracts and why they matter")
doc.add_paragraph("Boundary validation is defined once and inferred into TypeScript. Strict objects reject unknown fields, preventing accidental API widening and mass-assignment behavior.")
add_code(
    doc,
    "export const documentIngestJobSchema = z.strictObject({\n"
    "  schemaVersion: z.literal(1),\n"
    "  documentId: objectIdSchema,\n"
    "  processingVersion: z.number().int().positive(),\n"
    "  processingRunId: objectIdSchema,\n"
    "  correlationId: z.string().min(1).max(128),\n"
    "});",
)
add_table(
    doc,
    ["Code decision", "Why", "Failure prevented"],
    [
        ["schemaVersion literal", "Makes queue evolution explicit", "Old workers silently misreading new payloads"],
        ["documentId + processingVersion", "Pins work to one immutable source generation", "A retry overwriting a newer reprocess"],
        ["processingRunId", "Separates job attempt history from document state", "Ambiguous progress and cancellation"],
        ["correlationId", "Connects API request, queue job, worker logs, and audits", "Untraceable distributed failures"],
    ],
)
doc.add_paragraph("The generated-answer schema requires at least one citation per claim. That is only the first gate: the API then confirms each chunk belongs to the selected active version, the page matches stored spans, and the normalized excerpt occurs in the source text. If evidence is absent or validation fails, the result is an abstention—not a plausible-sounding answer.")

heading(doc, "6. Data ownership and tenant isolation")
add_table(
    doc,
    ["Record", "Owner key", "Critical invariants / indexes"],
    [
        ["WorkspaceMember", "workspaceId + userId", "Unique membership; role is owner/admin/member/viewer"],
        ["UploadIntent", "workspaceId", "Object key contains workspace; expiry and one completion timestamp"],
        ["Document", "workspaceId", "Unique upload intent; activeProcessingVersion; soft deletion"],
        ["ProcessingRun", "workspaceId + documentId", "Unique document/version; monotonic stageSequence; cancellation timestamp"],
        ["DocumentChunk", "workspaceId + documentId", "Unique document/version/stableChunkId; page spans; embedding version"],
        ["Conversation", "workspaceId", "Selected document IDs and optional collection fixed to tenant"],
        ["Message/Citation", "workspaceId + conversationId", "Client-request idempotency; evidence preserves document/chunk/version/page"],
    ],
)
add_callout(doc, "Authorization rule", "Never fetch by a public resource ID and authorize afterward. Resolve membership and include workspace, deletion, selected-source, and processing-version predicates inside the database or retrieval query itself.")

heading(doc, "7. Authentication flow")
add_code(
    doc,
    "register/login → Argon2id password check → HTTP-only access + opaque refresh cookies\n"
    "unsafe request → access JWT authentication → session lookup → CSRF HMAC verification\n"
    "expired access → one refresh rotation → old token revoked + replacement linked\n"
    "reused refresh → revoke all active user sessions → safe 409 response",
)
add_bullets(
    doc,
    [
        "Registration accepts exactly one identifier: normalized email or E.164 phone. Verification is deliberately omitted by product decision.",
        "Passwords are never logged or returned. Argon2id uses bounded parameters and login performs a fallback hash check to reduce user-enumeration timing differences.",
        "Access JWTs are short-lived and contain only subject, session ID, type, issuer, audience, issued-at, and expiration claims.",
        "Refresh tokens are stored only as a peppered digest. Rotation happens in a MongoDB transaction and detects replay.",
        "The browser stores only the CSRF value in session storage; authentication credentials remain HTTP-only cookies.",
    ],
)

heading(doc, "8. Direct upload and ingestion code path")
add_table(
    doc,
    ["Step", "Implementation", "Validation / recovery"],
    [
        ["1. Intent", "API validates filename, MIME, byte count, checksum, quota; reserves storage key", "Expired intents cannot complete"],
        ["2. Upload", "Browser computes SHA-256 and PUTs bytes to the signed storage URL", "Express memory is unaffected by PDF size"],
        ["3. Complete", "API HEADs object, checks size/type/hash metadata and `%PDF-` prefix", "Mismatch rejects before queueing"],
        ["4. Enqueue", "Deterministic job ID uses document/version", "Duplicate completion cannot duplicate ingestion"],
        ["5. Download", "Worker pipelines object stream through byte counter and SHA-256", "Temp byte cap; checksum; guaranteed cleanup"],
        ["6. Extract", "pdfjs reads pages in order and normalizes text", "Page cap and unsupported/corrupt PDF rejection"],
        ["7. Chunk", "Per-page overlapping windows with deterministic content hashes", "Stable reprocessing and exact page lineage"],
        ["8. Embed", "Gemini requests are batched and each dimension is validated", "Provider failures retry exponentially"],
        ["9. Commit", "Version chunks are replaced idempotently; document becomes ready", "Stale version/cancel checks before each stage"],
    ],
)
add_code(
    doc,
    "const jobId = `${documentId}:${processingVersion}`;\n"
    "await ingestQueue.add('document.ingest', validatedPayload, {\n"
    "  jobId, attempts: 3, backoff: { type: 'exponential', delay: 2000 }\n"
    "});",
)
doc.add_paragraph("This deterministic ID is the first idempotency fence. The worker adds database fences: it rejects stale versions, advances only to a greater stage sequence, upserts/replaces version-owned chunks, and never marks a cancelled run ready.")

heading(doc, "9. Retrieval and answer validation")
add_bullets(
    doc,
    [
        "Create the question embedding with the configured embedding model and version.",
        "Run MongoDB Atlas Vector Search and MongoDB Search with workspace, selected document, active version, and deletion metadata filters inside each pipeline.",
        "Fuse vector and lexical ranks using reciprocal-rank fusion; keep only the configured top K; enforce a minimum evidence score.",
        "Budget context by characters/tokens and format each evidence block with immutable chunk ID, document name, and page number.",
        "Tell the model that document text is untrusted evidence, not instruction. Request a strict JSON claims structure.",
        "Validate every citation against retrieved chunks, page spans, active versions, and normalized source excerpt. Persist only validated claims.",
        "If evidence is insufficient, the AI provider is unavailable, or validation fails, store a safe abstention and expose no fabricated citation.",
    ],
)
add_callout(doc, "Production index prerequisite", "Local MongoDB exercises persistence and orchestration but not Atlas vector/search operators. Before production or staging RAG tests, create both indexes using the mappings in the architecture plan and run the reproducible retrieval/citation evaluation dataset.")

heading(doc, "10. API and SSE behavior")
api_rows = [
    ["Auth", "POST /auth/register · /login · /refresh · /logout; GET /auth/me", "Cookie session + CSRF envelope"],
    ["Upload", "POST /uploads/intents; POST /uploads/:id/complete", "Signed PUT details; queued document/run"],
    ["Documents", "GET collection/item; PATCH name; DELETE; reprocess; cancel; view URL", "Tenant-scoped metadata or accepted job"],
    ["Collections", "POST/GET/PATCH/DELETE /document-collections", "Reusable document grouping"],
    ["Conversations", "POST/GET/PATCH conversations; POST questions", "JSON or text/event-stream"],
    ["Citations", "GET /citations/:id/source", "Short-lived URL + verified page/excerpt"],
]
add_table(doc, ["Area", "Routes", "Result"], api_rows)
doc.add_paragraph("Every error follows `{ error: { code, message, requestId, details? } }`. Validation errors expose safe field paths, expected operational failures use stable codes, and unexpected exceptions are logged with the request ID but returned as a generic 500. SSE connections send headers immediately, publish typed named events, send keepalives, and release Redis subscriptions when the client disconnects.")

heading(doc, "11. Frontend composition and visible reflection")
if (ROOT / "docs" / "design" / "askpdf-logo-study.png").exists():
    doc.add_picture(str(ROOT / "docs" / "design" / "askpdf-logo-study.png"), width=Inches(2.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_table(
    doc,
    ["Component / state", "Implementation concept", "What the user sees"],
    [
        ["AuthPage", "Local form state + SessionProvider mutation", "Email/phone segmented control, required password, sign-in/register mode"],
        ["AppShell", "Protected router outlet", "Sidebar, responsive drawer, private-workspace indicator, theme control"],
        ["DocumentsPage", "TanStack Query + server-side search and upload/collection mutations", "Ruled document catalogue, processing counts, empty/loading/error/success states, selection"],
        ["Upload modal", "Web Crypto SHA-256 + signed storage PUT", "Secure upload progress without routing bytes through API"],
        ["Collection modal", "Validated API mutation", "Create named collection from selected documents or empty"],
        ["ChatPage", "Conversation queries and question mutation", "Selected sources, grounded answers, citations, insufficiency state"],
        ["Citation drawer", "Short-lived source URL", "Evidence excerpt and exact-page PDF access"],
        ["Theme/responsive CSS", "Tokenized CSS custom properties + reduced-motion media query", "Beige editorial workspace, dark mode, purpose-built mobile navigation"],
    ],
)

heading(doc, "12. Tests, security gates, and release evidence")
add_table(
    doc,
    ["Gate", "What it proves", "Command"],
    [
        ["Formatting", "Repository-consistent source and documentation", "pnpm format:check"],
        ["Lint", "No configured correctness or safety warnings", "pnpm lint"],
        ["Typecheck", "Every workspace compiles under strict TS without emission", "pnpm typecheck"],
        ["Unit/integration", "Contracts, models, RAG primitives, errors, React auth rendering", "pnpm test"],
        ["Browser", "Desktop Chromium and mobile WebKit auth/theme flow", "pnpm test:e2e"],
        ["Production build", "All packages and three applications emit deployable output", "pnpm build"],
        ["Infrastructure", "Mongo/Redis/MinIO initialization and health", "docker compose up -d --wait"],
        ["Readiness", "Live Mongo connection plus Redis ping and storage bucket access", "Invoke-RestMethod http://localhost:4000/health/ready"],
    ],
)
add_bullets(
    doc,
    [
        "CI repeats format, lint, typecheck, tests, build, and Playwright in clean GitHub runners.",
        "Before production: add dependency/SBOM and secret scans, malicious/corrupt/encrypted/oversized PDF fixtures, authorization matrix tests, Redis/Mongo/S3/Gemini outage tests, and load tests at configured limits.",
        "RAG release gates must include Retrieval Hit@K, MRR, citation precision/coverage, answer faithfulness, unsupported-claim rate, correct-refusal rate, ingestion success, P95 query latency, and P95 processing time.",
        "Rollback independently reverts web/API/worker artifacts; additive schema changes remain backward compatible; pause ingestion before incompatible worker changes; preserve object and database backups.",
    ],
)

heading(doc, "13. Configuration reference")
config_rows = [
    ["Ports/origins", "WEB_PORT, API_HOST, API_PORT, API_PUBLIC_URL, WEB_ORIGIN", "5174/4000 locally; exact CORS origin"],
    ["Mongo/Redis", "MONGODB_URI, REDIS_URL, REDIS_KEY_PREFIX", "Credentials and tenant-safe key namespace"],
    ["Storage", "STORAGE_ENDPOINT/REGION/BUCKET/access keys/force path style", "Private S3-compatible bucket"],
    ["Session", "ACCESS_TOKEN_SECRET, REFRESH_TOKEN_PEPPER, CSRF_SECRET, cookie and TTL fields", "Use independent ≥32-character production secrets"],
    ["Gemini", "GEMINI_API_KEY, generation/embedding models, dimension, timeout", "Key must be supplied by deployer"],
    ["File bounds", "MAX_UPLOAD_BYTES, MAX_DOCUMENT_PAGES, MAX_EXTRACTED_TEXT_CHARACTERS, MAX_TEMP_BYTES_PER_JOB", "100 MiB, 500 pages, 10M chars, 150 MiB proposed defaults"],
    ["Worker", "PDF_PAGE_BATCH_SIZE, EMBEDDING_BATCH_SIZE, WORKER_CONCURRENCY, job attempts/backoff", "10 pages, 32 embeddings, 2 jobs, 3 attempts/2s proposed defaults"],
    ["Retrieval", "chunk target/overlap, top K, evidence score, max context", "2600/400 chars, top 8, 0.35, 28k chars proposed defaults"],
]
add_table(doc, ["Group", "Variables", "Meaning / current proposed default"], config_rows)
add_callout(doc, "Secrets", "`.env` is ignored and must never be committed. `.env.example` contains only local development values and explicit replacement strings. Production keys belong in the deployment secret manager and should be rotated after any suspected exposure.", "FFF5E8")

heading(doc, "14. Master command sheet")
add_code(
    doc,
    "# one-time workstation setup\n"
    "corepack enable\n"
    "corepack prepare pnpm@11.9.0 --activate\n"
    "Copy-Item .env.example .env\n"
    "pnpm install --frozen-lockfile\n"
    "pnpm exec playwright install chromium webkit\n\n"
    "# infrastructure + application\n"
    "docker compose up -d --wait\n"
    "pnpm dev\n\n"
    "# validation\n"
    "pnpm format:check\n"
    "pnpm lint\n"
    "pnpm typecheck\n"
    "pnpm test\n"
    "pnpm test:e2e\n"
    "pnpm build\n"
    "pnpm verify\n\n"
    "# diagnostics\n"
    "docker compose ps\n"
    "docker compose logs --tail=100 mongo redis minio\n"
    "Invoke-RestMethod http://localhost:4000/health/live\n"
    "Invoke-RestMethod http://localhost:4000/health/ready",
)
add_callout(doc, "Safe Docker reset", "`docker compose down` stops the stack and preserves named volumes. Add `-v` only when intentionally deleting local documents, queues, and metadata; that destructive variant is not part of routine troubleshooting.")

heading(doc, "15. Operational limits and graceful degradation")
add_table(
    doc,
    ["Dependency / pressure", "Behavior", "Operator action"],
    [
        ["Gemini unavailable", "No ungrounded fallback; ingestion retries; questions return dependency-unavailable/abstention", "Inspect provider quota/key/status; resume deterministic jobs"],
        ["Redis unavailable", "Queue/progress/readiness fail; API does not claim ready", "Restore Redis; BullMQ resumes persisted jobs"],
        ["MongoDB unavailable", "Readiness fails; no metadata mutation", "Restore replica set; verify transactions and indexes"],
        ["S3/MinIO unavailable", "Signed/object metadata operations fail safely; readiness fails", "Restore bucket access; retry unchanged intent/job"],
        ["Oversized/corrupt PDF", "Rejected before or during processing with bounded resource use", "User replaces source; no manual database repair"],
        ["Provider rate pressure", "Bounded worker concurrency, embedding batches, exponential retry", "Tune per-workspace quotas/concurrency from telemetry"],
        ["Worker crash", "Job retry uses deterministic ID/version and guarded stage checkpoints", "Restart worker; stale/cancelled work is rejected"],
    ],
    trailing_space=False,
    font_size=7,
)

heading(doc, "16. Decisions still owned by the deployer")
add_bullets(
    doc,
    [
        "Provide the production Gemini API key and confirm the Gemini region/data-governance policy.",
        "Provide the Git remote URL and repository access if this local commit must be pushed.",
        "Choose production MongoDB Atlas, Redis, and S3-compatible endpoints and secrets.",
        "Approve v2 scope separately before implementing learning activities.",
        "Set product quotas and evidence threshold after staging load tests and the first labeled evaluation run.",
    ],
)
doc.core_properties.title = "AskPDF Implementation and Code Guide"
doc.core_properties.subject = "Architecture, build order, code map, security, testing, and operations"
doc.core_properties.author = "AskPDF Engineering"
doc.core_properties.keywords = "AskPDF, RAG, TypeScript, Gemini, MongoDB, BullMQ, citations"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
